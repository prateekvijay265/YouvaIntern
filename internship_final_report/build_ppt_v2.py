"""
build_ppt_v2.py
===============
Generates a 10-slide professional PowerPoint presentation for the AI Internship.
Uses a modern tech-savvy color palette while maintaining the angular, ribbon-based 
layout aesthetic from the user's reference slides.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Paths ─────────────────────────────────────────────────────────────────────
BASE_DIR = r"c:\Users\prate\Desktop\youva"
REPORT_DIR = os.path.join(BASE_DIR, "internship_final_report")
DP_DIR = os.path.join(BASE_DIR, "data processing")
ML_DIR = os.path.join(BASE_DIR, "ml_model_development", "results", "figures")
XAI_DIR = os.path.join(BASE_DIR, "explainable_ai")
DEPLOY_DIR = os.path.join(BASE_DIR, "ai_deployment_monitoring")
ETHICS_DIR = os.path.join(BASE_DIR, "ai_ethics_audit")

# ── Tech-Savvy Color Palette ──────────────────────────────────────────────────
# Background & Text
BG_COLOR       = RGBColor(0xF1, 0xF5, 0xF9)  # Light cool gray (modern background)
TEXT_DARK      = RGBColor(0x0F, 0x17, 0x2A)  # Very dark slate (for main headings)
TEXT_SUB       = RGBColor(0x33, 0x41, 0x55)  # Slate gray (for body text)

# Ribbon Colors
RIBBON_MAIN    = RGBColor(0x25, 0x63, 0xEB)  # Vibrant Tech Blue
RIBBON_FOLD    = RGBColor(0x1E, 0x3A, 0x8A)  # Deep Navy (for shadows/folds)
RIBBON_ACCENT  = RGBColor(0x06, 0xB6, 0xD4)  # Electric Cyan
RIBBON_GREY    = RGBColor(0x94, 0xA3, 0xB8)  # Muted slate ribbon

WHITE          = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── Helpers ───────────────────────────────────────────────────────────────────

def set_slide_bg(slide, color=BG_COLOR):
    """Set solid background color for a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, left, top, width, height, fill_color, adj=None):
    """Add a generic shape with solid fill and no border."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    if adj is not None:
        shape.adjustments[0] = adj
    return shape


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=TEXT_DARK, bold=False, alignment=PP_ALIGN.LEFT,
                font_name='Segoe UI'):
    """Add a styled text box."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_frame(slide, left, top, width, height, items, font_size=16,
                     color=TEXT_SUB):
    """Add a text frame with bulleted items."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = 'Segoe UI'
        p.space_after = Pt(10)
        p.level = 0
    return txBox


def add_image_safe(slide, path, left, top, width=None, height=None):
    """Add image if it exists."""
    if os.path.exists(path):
        kwargs = {}
        if width: kwargs['width'] = width
        if height: kwargs['height'] = height
        slide.shapes.add_picture(path, left, top, **kwargs)
        return True
    return False


def apply_theme_ribbons(slide, variant="default"):
    """Applies the angular ribbon aesthetic to a slide."""
    set_slide_bg(slide)
    
    if variant == "title":
        # Large angular block on right
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(7), Inches(-1), Inches(8), Inches(9.5), RIBBON_GREY, adj=0.2)
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(7.5), Inches(-1), Inches(8), Inches(9.5), WHITE, adj=0.2)
        
        # Ribbons crossing bottom left
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(-1), Inches(5), Inches(8), Inches(1.5), RIBBON_MAIN, adj=0.3)
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(1), Inches(5.8), Inches(7), Inches(1.2), RIBBON_FOLD, adj=0.3)
        
    elif variant == "default":
        # Top ribbon
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(8), Inches(-0.5), Inches(6), Inches(1.5), RIBBON_ACCENT, adj=0.3)
        # Mid-left ribbon
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(-1), Inches(3.5), Inches(4), Inches(1.2), RIBBON_MAIN, adj=0.3)
        # Bottom right angular accent
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(9), Inches(6.5), Inches(5), Inches(1.5), RIBBON_FOLD, adj=0.3)

    elif variant == "content_heavy":
        # Left vertical ribbon
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(-1), Inches(-1), Inches(2.5), Inches(9.5), RIBBON_FOLD, adj=0.1)
        add_shape(slide, MSO_SHAPE.PARALLELOGRAM, Inches(1), Inches(5), Inches(13), Inches(1.5), RIBBON_MAIN, adj=0.2)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE GENERATION
# ══════════════════════════════════════════════════════════════════════════════

def generate_ppt():
    print("=" * 50)
    print("  GENERATING TECH-SAVVY PRESENTATION")
    print("=" * 50)
    
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]
    
    # ── SLIDE 1: TITLE ────────────────────────────────────────────────────────
    print("[1/10] Title slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "title")
    
    add_textbox(slide, Inches(1), Inches(1.5), Inches(6), Inches(1),
                "AI INTERNSHIP\nPROJECT REPORT", font_size=44, color=TEXT_DARK, bold=True)
    
    add_textbox(slide, Inches(1), Inches(3.0), Inches(6), Inches(0.5),
                "6-Week Intensive AI/ML Program", font_size=20, color=RIBBON_MAIN, bold=True)
    
    add_textbox(slide, Inches(9.5), Inches(3.5), Inches(3), Inches(1),
                "Presented By:\n[Student Name]", font_size=16, color=TEXT_DARK, alignment=PP_ALIGN.RIGHT)
    add_textbox(slide, Inches(9.5), Inches(4.5), Inches(3), Inches(0.5),
                "2025 - 2026", font_size=24, color=RIBBON_MAIN, bold=True, alignment=PP_ALIGN.RIGHT)
    
    # ── SLIDE 2: OVERVIEW ─────────────────────────────────────────────────────
    print("[2/10] Overview slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "default")
    
    add_textbox(slide, Inches(1), Inches(0.8), Inches(6), Inches(1),
                "PROGRAM OVERVIEW", font_size=36, color=TEXT_DARK, bold=True)
    add_textbox(slide, Inches(1), Inches(1.7), Inches(6), Inches(1),
                "A comprehensive journey through the AI lifecycle, from planning to production deployment.",
                font_size=18, color=RIBBON_MAIN)
    
    add_textbox(slide, Inches(1), Inches(2.8), Inches(6), Inches(0.5),
                "KEY PHASES", font_size=20, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(1), Inches(3.4), Inches(6), Inches(3),
        [
            "Week 1: AI Project Planning & Feasibility",
            "Week 2: Data Preprocessing & Feature Engineering",
            "Week 3: ML Model Building & Tuning",
            "Week 4: Explainable AI & Interpretability",
            "Week 5: API Deployment & Monitoring",
            "Week 6: AI Ethics & Bias Analysis"
        ], font_size=16)
    
    # ── SLIDE 3: WEEK 1 ───────────────────────────────────────────────────────
    print("[3/10] Week 1 slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "content_heavy")
    
    add_textbox(slide, Inches(2), Inches(0.8), Inches(8), Inches(0.8),
                "WEEK 1: PROJECT PLANNING", font_size=32, color=TEXT_DARK, bold=True)
    
    add_textbox(slide, Inches(2), Inches(1.8), Inches(9), Inches(0.5),
                "DELIVERABLES & STRATEGY", font_size=18, color=WHITE, bold=True)
    
    add_bullet_frame(slide, Inches(2), Inches(2.4), Inches(9), Inches(2),
        [
            "Formulated hypothesis for healthcare predictive analytics.",
            "Selected optimal technology stack: Python, scikit-learn, pandas.",
            "Designed project timeline with milestone tracking.",
            "Identified risks and proposed mitigation strategies."
        ], font_size=16)
    
    # ── SLIDE 4: WEEK 2 ───────────────────────────────────────────────────────
    print("[4/10] Week 2 slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "default")
    
    add_textbox(slide, Inches(1), Inches(0.8), Inches(8), Inches(0.8),
                "WEEK 2: DATA PREPROCESSING", font_size=32, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(1), Inches(1.8), Inches(6), Inches(4),
        [
            "Dataset: Titanic (891 samples)",
            "Applied KNN Imputation for missing Age values (19.9%).",
            "Executed IQR-based outlier detection with Winsorization.",
            "Engineered features: FamilySize, IsAlone, Title, FareBin.",
            "Built a reproducible scikit-learn Pipeline."
        ], font_size=16)
    
    img1 = os.path.join(DP_DIR, "correlation_matrix.png")
    add_image_safe(slide, img1, Inches(7.5), Inches(1.8), width=Inches(5))
    
    # ── SLIDE 5: WEEK 3 ───────────────────────────────────────────────────────
    print("[5/10] Week 3 slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "content_heavy")
    
    add_textbox(slide, Inches(2), Inches(0.8), Inches(8), Inches(0.8),
                "WEEK 3: ML MODEL BUILDING", font_size=32, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(2), Inches(1.8), Inches(6), Inches(3),
        [
            "Algorithm: Random Forest Classifier",
            "Dataset: Breast Cancer Wisconsin",
            "Tuned 162 hyperparameter combinations via GridSearchCV.",
            "Evaluated using 5-fold Stratified Cross-Validation.",
            "Achieved F1-Score of 0.99 and ROC-AUC of 0.999."
        ], font_size=16)
    
    img2 = os.path.join(ML_DIR, "roc_curves.png")
    add_image_safe(slide, img2, Inches(8.5), Inches(1.5), width=Inches(4))
    
    # ── SLIDE 6: WEEK 4 ───────────────────────────────────────────────────────
    print("[6/10] Week 4 slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "default")
    
    add_textbox(slide, Inches(1), Inches(0.8), Inches(8), Inches(0.8),
                "WEEK 4: EXPLAINABLE AI (XAI)", font_size=32, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(1), Inches(1.8), Inches(5.5), Inches(3),
        [
            "Implemented SHAP (SHapley Additive exPlanations) for global feature importance.",
            "Utilized LIME for local surrogate explanations of individual predictions.",
            "Established clinical trust through transparent decision breakdowns."
        ], font_size=16)
    
    img3 = os.path.join(XAI_DIR, "fig3_shap_summary.png")
    add_image_safe(slide, img3, Inches(7), Inches(1.5), width=Inches(5.5))
    
    # ── SLIDE 7: WEEK 5 (Part 1) ──────────────────────────────────────────────
    print("[7/10] Week 5 (Part 1) slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "content_heavy")
    
    add_textbox(slide, Inches(2), Inches(0.8), Inches(10), Inches(0.8),
                "WEEK 5: DEPLOYMENT ARCHITECTURE", font_size=32, color=TEXT_DARK, bold=True)
    add_textbox(slide, Inches(2), Inches(1.5), Inches(10), Inches(0.4),
                "★ HIGHEST SCORE: 95/100", font_size=16, color=WHITE, bold=True)
    
    add_bullet_frame(slide, Inches(2), Inches(2.2), Inches(10), Inches(3),
        [
            "Built a production-ready REST API using FastAPI.",
            "Implemented robust Pydantic schemas for request validation.",
            "Containerized the service using a multi-stage Docker build.",
            "Configured Docker Compose with resource limits and health checks."
        ], font_size=16)
    
    # ── SLIDE 8: WEEK 5 (Part 2) ──────────────────────────────────────────────
    print("[8/10] Week 5 (Part 2) slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "default")
    
    add_textbox(slide, Inches(1), Inches(0.8), Inches(10), Inches(0.8),
                "WEEK 5: MONITORING & TESTING", font_size=32, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(1), Inches(1.8), Inches(11), Inches(3),
        [
            "Thread-Safe Metrics: Tracked request counts, latencies (P95), and error rates.",
            "Structured Logging: Used python-json-logger with rotating file handlers.",
            "Automated Testing: Developed a comprehensive test suite of 15 pytest tests.",
            "Production Features: Added Request-ID middleware and global exception handling."
        ], font_size=16)
    
    # ── SLIDE 9: WEEK 6 ───────────────────────────────────────────────────────
    print("[9/10] Week 6 slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "content_heavy")
    
    add_textbox(slide, Inches(2), Inches(0.8), Inches(10), Inches(0.8),
                "WEEK 6: AI ETHICS & BIAS", font_size=32, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(2), Inches(1.8), Inches(6), Inches(4),
        [
            "Audited the UCI Adult Income dataset for demographic biases.",
            "Evaluated metrics: Disparate Impact, Equalized Odds, Predictive Parity.",
            "Identified intersectional bias affecting non-White females.",
            "Proposed remediation: Resampling and adversarial debiasing."
        ], font_size=16)
    
    img4 = os.path.join(ETHICS_DIR, "fig2_disparate_impact.png")
    add_image_safe(slide, img4, Inches(8.5), Inches(1.5), width=Inches(4))
    
    # ── SLIDE 10: CONCLUSION ──────────────────────────────────────────────────
    print("[10/10] Conclusion slide...")
    slide = prs.slides.add_slide(blank_layout)
    apply_theme_ribbons(slide, "title")  # Reuse title layout for strong finish
    
    add_textbox(slide, Inches(1), Inches(1.5), Inches(6), Inches(1),
                "CONCLUSION & FUTURE", font_size=36, color=TEXT_DARK, bold=True)
    
    add_bullet_frame(slide, Inches(1), Inches(2.5), Inches(6), Inches(3),
        [
            "End-to-end understanding of the AI lifecycle achieved.",
            "Deployment engineering is as critical as model accuracy.",
            "Ethics and fairness must be integrated early.",
            "Future: Implement CI/CD and cloud deployment (AWS/GCP)."
        ], font_size=16)
    
    # ══════════════════════════════════════════════════════════════════════
    # SAVE
    # ══════════════════════════════════════════════════════════════════════
    pptx_path = os.path.join(REPORT_DIR, "Internship_Presentation_V2.pptx")
    prs.save(pptx_path)
    print(f"\n[OK] Tech-Savvy Presentation saved: {pptx_path}")
    print(f"     Total slides: {len(prs.slides)}")
    print("\n" + "=" * 50)
    print("  PRESENTATION COMPLETE")
    print("=" * 50)


if __name__ == "__main__":
    generate_ppt()
