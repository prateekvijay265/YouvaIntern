"""
build_ppt_v5.py — Image-background approach for premium look.
Uses generated background images + clean text overlays.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE   = r"c:\Users\prate\Desktop\youva"
OUT    = os.path.join(BASE, "internship_final_report")
IMG_DP = os.path.join(BASE, "data processing")
IMG_ML = os.path.join(BASE, "ml_model_development", "results", "figures")
IMG_XAI = os.path.join(BASE, "explainable_ai")
IMG_ETH = os.path.join(BASE, "ai_ethics_audit")
ART    = r"C:\Users\prate\.gemini\antigravity\brain\90c607a8-2b17-4c52-9d23-7a91d123a082"

# Background images
BG_TITLE   = os.path.join(ART, "slide_bg_title_1786209103973.jpg")
BG_CONTENT = os.path.join(ART, "slide_bg_content_left_1786209120429.jpg")
BG_FULL    = os.path.join(ART, "slide_bg_full_1786209136708.jpg")
BG_DARK    = os.path.join(ART, "slide_bg_dark_1786209157393.jpg")

NAVY   = RGBColor(0x0F, 0x24, 0x4E)
TEAL   = RGBColor(0x00, 0xB4, 0xD8)
AMBER  = RGBColor(0xF5, 0xA6, 0x23)
CORAL  = RGBColor(0xEF, 0x47, 0x6F)
MINT   = RGBColor(0x2D, 0xD4, 0xBF)
SLATE  = RGBColor(0x47, 0x55, 0x69)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY  = RGBColor(0xCB, 0xD5, 0xE1)

SW = Inches(13.333)
SH = Inches(7.5)

def set_bg_img(slide, img_path):
    """Set a full-slide background image."""
    slide.shapes.add_picture(img_path, Emu(0), Emu(0), SW, SH)

def txt(slide, l, t, w, h, text, sz=18, c=NAVY, bold=False, align=PP_ALIGN.LEFT, font='Segoe UI'):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text; p.font.size = Pt(sz); p.font.color.rgb = c
    p.font.bold = bold; p.font.name = font; p.alignment = align
    return tb

def multi(slide, l, t, w, h, lines, sz=14, c=SLATE, font='Segoe UI', spacing=Pt(8)):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.size = Pt(sz); p.font.color.rgb = c
        p.font.name = font; p.space_after = spacing
    return tb

def rect(slide, l, t, w, h, c, r=None):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = c
    sh.line.fill.background()
    if r: sh.adjustments[0] = r
    return sh

def badge(slide, l, t, text, c=AMBER):
    sh = rect(slide, l, t, Inches(1.4), Inches(0.45), c, r=0.2)
    tf = sh.text_frame; p = tf.paragraphs[0]
    p.text = text; p.font.size = Pt(13); p.font.color.rgb = WHITE
    p.font.bold = True; p.font.name = 'Segoe UI'; p.alignment = PP_ALIGN.CENTER

def add_img(slide, path, l, t, width=None, height=None):
    if not os.path.exists(path): return
    kw = {}
    if width: kw['width'] = width
    if height: kw['height'] = height
    slide.shapes.add_picture(path, l, t, **kw)


def build():
    print("BUILDING v5 (image backgrounds)")
    print("=" * 40)

    prs = Presentation()
    prs.slide_width = SW; prs.slide_height = SH
    BL = prs.slide_layouts[6]

    # ── SLIDE 1: TITLE ────────────────────────────────────────────────────
    print("[1/10] Title")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_TITLE)

    txt(s, Inches(0.9), Inches(1.2), Inches(5.5), Inches(0.5),
        "COMPREHENSIVE", 18, TEAL, True)
    txt(s, Inches(0.9), Inches(1.7), Inches(5.5), Inches(1.5),
        "AI INTERNSHIP\nPROJECT REPORT", 44, NAVY, True)
    txt(s, Inches(0.9), Inches(3.8), Inches(5.5), Inches(0.5),
        "A 6-Week Deep Dive into AI -- From Planning to Production", 15, SLATE)

    txt(s, Inches(8.5), Inches(2.5), Inches(4), Inches(0.4),
        "Presented By", 13, LGRAY)
    txt(s, Inches(8.5), Inches(2.9), Inches(4), Inches(0.6),
        "Prateek Vijay", 26, WHITE, True)
    txt(s, Inches(8.5), Inches(3.6), Inches(4), Inches(0.4),
        "May -- June 2026", 15, TEAL, True)
    txt(s, Inches(8.5), Inches(4.1), Inches(4), Inches(0.4),
        "AI/ML Internship Program", 13, LGRAY)

    # ── SLIDE 2: OVERVIEW ─────────────────────────────────────────────────
    print("[2/10] Overview")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_FULL)

    txt(s, Inches(0.9), Inches(0.5), Inches(10), Inches(0.7),
        "WHAT THIS INTERNSHIP COVERED", 32, NAVY, True)
    txt(s, Inches(0.9), Inches(1.2), Inches(10), Inches(0.4),
        "Six weeks. Six real projects. One complete AI lifecycle.", 17, TEAL, True)

    weeks = [
        ("01", "Project Planning",    "Scoped a healthcare predictive analytics\nsystem with risk assessment & timelines", TEAL,  "65"),
        ("02", "Data Preprocessing",  "Cleaned the Titanic dataset with KNN\nimputation & built a reusable pipeline",     MINT,  "85"),
        ("03", "Model Building",      "Tuned a Random Forest across 162\ncombinations on breast cancer data",             AMBER, "45"),
        ("04", "Explainable AI",      "Made the model transparent with\nSHAP & LIME interpretability",                    CORAL, "55"),
        ("05", "Deployment",          "Shipped a FastAPI + Docker service\nwith monitoring & 15 automated tests",          TEAL,  "95"),
        ("06", "Ethics & Bias",       "Audited 32k records for fairness\nacross sex & race dimensions",                   NAVY,  "70"),
    ]
    for i, (num, title, desc, color, score) in enumerate(weeks):
        col = i % 3; row = i // 3
        x = Inches(0.5 + col * 4.15); y = Inches(1.9 + row * 2.7)

        rect(s, x, y, Inches(3.9), Inches(2.4), WHITE, r=0.02)
        rect(s, x, y, Inches(3.9), Pt(5), color)
        # Number
        n = rect(s, x + Inches(0.2), y + Inches(0.25), Inches(0.55), Inches(0.55), color, r=0.5)
        tf = n.text_frame; p = tf.paragraphs[0]
        p.text = num; p.font.size = Pt(15); p.font.color.rgb = WHITE
        p.font.bold = True; p.font.name = 'Segoe UI'; p.alignment = PP_ALIGN.CENTER
        txt(s, x + Inches(0.9), y + Inches(0.25), Inches(2.8), Inches(0.45), title, 14, NAVY, True)
        txt(s, x + Inches(0.2), y + Inches(0.9), Inches(3.4), Inches(0.9), desc, 11, SLATE)
        badge(s, x + Inches(2.3), y + Inches(1.8), f"{score}/100", color)

    # ── SLIDE 3: WEEK 1 ───────────────────────────────────────────────────
    print("[3/10] Week 1")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_FULL)

    txt(s, Inches(0.9), Inches(0.3), Inches(3), Inches(0.3), "WEEK 1", 13, TEAL, True)
    badge(s, Inches(11.5), Inches(0.3), "65/100", AMBER)
    txt(s, Inches(0.9), Inches(0.7), Inches(10), Inches(0.6),
        "PROJECT PLANNING & DOCUMENTATION", 28, NAVY, True)

    txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.7),
        "I started by asking: \"What if we could predict patient readmission risk "
        "before discharge?\" That single hypothesis shaped every decision -- from "
        "dataset selection to model evaluation criteria.", 14, SLATE)

    # Two cards
    rect(s, Inches(0.5), Inches(2.6), Inches(6.0), Inches(4.4), WHITE, r=0.02)
    rect(s, Inches(0.5), Inches(2.6), Inches(6.0), Pt(5), TEAL)
    txt(s, Inches(0.8), Inches(2.8), Inches(5.5), Inches(0.4), "WHAT I DELIVERED", 15, TEAL, True)
    multi(s, Inches(0.8), Inches(3.3), Inches(5.5), Inches(3.5), [
        "  A full project plan targeting healthcare readmissions",
        "  Hypothesis: clinical + demographic features can",
        "      predict 30-day readmission with AUC > 0.80",
        "  Tech stack rationale -- Python, scikit-learn, pandas",
        "  Risk register covering data quality, overfitting,",
        "      compute constraints, and patient privacy",
        "  Project timeline with phased milestones",
    ], 12, SLATE)

    rect(s, Inches(6.8), Inches(2.6), Inches(6.0), Inches(4.4), WHITE, r=0.02)
    rect(s, Inches(6.8), Inches(2.6), Inches(6.0), Pt(5), CORAL)
    txt(s, Inches(7.1), Inches(2.8), Inches(5.5), Inches(0.4), "WHAT I LEARNED", 15, CORAL, True)
    multi(s, Inches(7.1), Inches(3.3), Inches(5.5), Inches(3.5), [
        "  Documentation depth matters -- the evaluator",
        "      wanted more on ethics and data governance",
        "  A good hypothesis drives the entire project scope;",
        "      without it, everything else is unfocused",
        "  Planning isn't just paperwork -- it saves weeks",
        "      of rework later in the project lifecycle",
        "  Score deducted for not meeting the word count",
    ], 12, SLATE)

    # ── SLIDE 4: WEEK 2 ───────────────────────────────────────────────────
    print("[4/10] Week 2")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_CONTENT)

    txt(s, Inches(0.9), Inches(0.3), Inches(3), Inches(0.3), "WEEK 2", 13, MINT, True)
    badge(s, Inches(5.5), Inches(0.3), "85/100", MINT)
    txt(s, Inches(0.9), Inches(0.7), Inches(6), Inches(0.6),
        "DATA PREPROCESSING", 28, NAVY, True)
    txt(s, Inches(0.9), Inches(1.4), Inches(6), Inches(0.7),
        "The Titanic dataset has 891 passengers and a deliberate "
        "mess of missing data -- exactly the kind of real-world "
        "chaos you need to learn to handle.", 13, SLATE)

    rect(s, Inches(0.5), Inches(2.4), Inches(6.3), Inches(4.7), WHITE, r=0.02)
    rect(s, Inches(0.5), Inches(2.4), Inches(6.3), Pt(5), MINT)
    txt(s, Inches(0.8), Inches(2.6), Inches(5.8), Inches(0.4), "KEY TECHNIQUES", 15, MINT, True)
    multi(s, Inches(0.8), Inches(3.1), Inches(5.8), Inches(3.8), [
        "  KNN Imputation for Age -- smarter than just",
        "  plugging in the mean; uses feature similarity",
        "",
        "  IQR method with Winsorization -- tamed fare",
        "  outliers without throwing away valuable data",
        "",
        "  Engineered FamilySize, IsAlone, Title, FareBin",
        "  -- features that actually explain who survived",
        "",
        "  Wrapped everything in a sklearn Pipeline for",
        "  perfect reproducibility on new data",
    ], 12, SLATE)

    add_img(s, os.path.join(IMG_DP, "correlation_matrix.png"),
            Inches(7.2), Inches(0.3), width=Inches(5.6))
    add_img(s, os.path.join(IMG_DP, "feature_importance.png"),
            Inches(7.2), Inches(3.8), width=Inches(5.6))

    # ── SLIDE 5: WEEK 3 ───────────────────────────────────────────────────
    print("[5/10] Week 3")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_CONTENT)

    txt(s, Inches(0.9), Inches(0.3), Inches(3), Inches(0.3), "WEEK 3", 13, AMBER, True)
    badge(s, Inches(5.5), Inches(0.3), "45/100", RGBColor(0x94,0xA3,0xB8))
    txt(s, Inches(0.9), Inches(0.7), Inches(6), Inches(0.6),
        "BUILDING & TUNING THE MODEL", 26, NAVY, True)
    txt(s, Inches(0.9), Inches(1.4), Inches(6), Inches(0.7),
        "Trained a Random Forest on 569 breast cancer samples -- "
        "30 features from cell-nucleus images. Goal: distinguish "
        "malignant from benign tumours.", 13, SLATE)

    # Stat cards in 2x2 grid
    stats = [("162", "Hyperparameter\nCombinations", TEAL),
             ("810", "Total CV\nFits", AMBER),
             ("0.99", "Best F1\nScore", MINT),
             ("0.999", "Best ROC\nAUC", CORAL)]
    for i, (val, label, color) in enumerate(stats):
        col = i % 2; row = i // 2
        x = Inches(0.5 + col * 3.3); y = Inches(2.5 + row * 2.3)
        rect(s, x, y, Inches(3.0), Inches(2.0), WHITE, r=0.02)
        rect(s, x, y, Inches(3.0), Pt(5), color)
        txt(s, x + Inches(0.2), y + Inches(0.3), Inches(2.6), Inches(0.7),
            val, 30, color, True, PP_ALIGN.CENTER)
        txt(s, x + Inches(0.2), y + Inches(1.1), Inches(2.6), Inches(0.6),
            label, 12, SLATE, False, PP_ALIGN.CENTER)

    add_img(s, os.path.join(IMG_ML, "roc_curves.png"),
            Inches(7.2), Inches(0.3), width=Inches(5.6))
    add_img(s, os.path.join(IMG_ML, "confusion_matrices.png"),
            Inches(7.2), Inches(3.8), width=Inches(5.6))

    # ── SLIDE 6: WEEK 4 ───────────────────────────────────────────────────
    print("[6/10] Week 4")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_CONTENT)

    txt(s, Inches(0.9), Inches(0.3), Inches(3), Inches(0.3), "WEEK 4", 13, CORAL, True)
    badge(s, Inches(5.5), Inches(0.3), "55/100", RGBColor(0x94,0xA3,0xB8))
    txt(s, Inches(0.9), Inches(0.7), Inches(6), Inches(0.6),
        "EXPLAINABLE AI", 28, NAVY, True)
    txt(s, Inches(0.9), Inches(1.4), Inches(6), Inches(0.5),
        "\"Why did the model say malignant?\" If you can't "
        "answer that, no doctor will trust your prediction.", 13, SLATE)

    # SHAP card (dark)
    rect(s, Inches(0.5), Inches(2.2), Inches(6.3), Inches(2.3), NAVY, r=0.02)
    txt(s, Inches(0.8), Inches(2.4), Inches(5.8), Inches(0.4),
        "SHAP -- Global & Local Explanations", 15, TEAL, True)
    txt(s, Inches(0.8), Inches(2.9), Inches(5.8), Inches(1.4),
        "TreeExplainer revealed worst concave points and worst "
        "perimeter as top drivers -- aligning with what "
        "oncologists know about cell shape irregularity. "
        "Waterfall plots broke down individual predictions "
        "showing exactly what pushed each decision.", 12, LGRAY)

    # LIME card (dark)
    rect(s, Inches(0.5), Inches(4.8), Inches(6.3), Inches(2.3), NAVY, r=0.02)
    txt(s, Inches(0.8), Inches(5.0), Inches(5.8), Inches(0.4),
        "LIME -- Case-by-Case Transparency", 15, AMBER, True)
    txt(s, Inches(0.8), Inches(5.5), Inches(5.8), Inches(1.4),
        "Generated local surrogate explanations for individual "
        "predictions -- one benign, one malignant. Each shows "
        "exactly which measurements contributed and by how much. "
        "Complementary to SHAP for clinical trust.", 12, LGRAY)

    add_img(s, os.path.join(IMG_XAI, "fig3_shap_summary.png"),
            Inches(7.2), Inches(0.3), width=Inches(5.6))
    add_img(s, os.path.join(IMG_XAI, "fig6_lime_benign.png"),
            Inches(7.2), Inches(3.8), width=Inches(5.6))

    # ── SLIDE 7: WEEK 5 pt1 ──────────────────────────────────────────────
    print("[7/10] Week 5 pt1")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_DARK)

    txt(s, Inches(0.9), Inches(0.3), Inches(6), Inches(0.3),
        "WEEK 5  --  HIGHEST SCORE", 13, TEAL, True)
    badge(s, Inches(11.5), Inches(0.3), "95/100", TEAL)
    txt(s, Inches(0.9), Inches(0.7), Inches(10), Inches(0.6),
        "DEPLOYMENT & MONITORING", 30, WHITE, True)
    txt(s, Inches(0.9), Inches(1.4), Inches(11), Inches(0.6),
        "I didn't just train a model -- I shipped it as a real service that handles "
        "requests, logs everything, monitors itself, and restarts if something breaks.", 14, LGRAY)

    # 4 cards on dark background
    cards = [
        ("FastAPI",    "7 REST endpoints\nPydantic validation\nSwagger UI & CORS\nRequest-ID middleware", TEAL),
        ("Docker",     "Multi-stage build\nNon-root user\n512 MB memory cap\nAuto-restart policy",       RGBColor(0x38,0x4F,0x7C)),
        ("Monitoring", "Thread-safe metrics\nLatency P95 tracking\nError rate per-class\nRolling 1k window", AMBER),
        ("Testing",    "15 pytest tests\nAll endpoints covered\nEdge case validation\nMetric accumulation",  CORAL),
    ]
    for i, (title, desc, color) in enumerate(cards):
        x = Inches(0.5 + i * 3.15)
        rect(s, x, Inches(2.3), Inches(2.9), Inches(2.8), RGBColor(0x1A,0x32,0x5E), r=0.02)
        rect(s, x, Inches(2.3), Inches(2.9), Inches(0.6), color, r=0.02)
        txt(s, x + Inches(0.1), Inches(2.35), Inches(2.7), Inches(0.5),
            title, 16, WHITE, True, PP_ALIGN.CENTER)
        txt(s, x + Inches(0.2), Inches(3.1), Inches(2.5), Inches(1.8),
            desc, 12, LGRAY)

    # Endpoint bar
    rect(s, Inches(0.5), Inches(5.5), Inches(12.3), Inches(1.5), RGBColor(0x08,0x14,0x30), r=0.02)
    txt(s, Inches(0.8), Inches(5.6), Inches(11.7), Inches(0.3),
        "API ENDPOINTS", 12, TEAL, True)
    multi(s, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.8), [
        "GET /health         GET /metrics        GET /model/info",
        "POST /predict       POST /predict/batch     GET /docs       GET /redoc",
    ], 13, LGRAY, 'Consolas', Pt(4))

    # ── SLIDE 8: WEEK 5 pt2 ──────────────────────────────────────────────
    print("[8/10] Week 5 pt2")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_CONTENT)

    txt(s, Inches(0.9), Inches(0.3), Inches(6), Inches(0.3),
        "WEEK 5 -- UNDER THE HOOD", 13, TEAL, True)
    txt(s, Inches(0.9), Inches(0.7), Inches(6), Inches(0.6),
        "HOW IT ACTUALLY WORKS", 26, NAVY, True)

    # Docker card
    rect(s, Inches(0.5), Inches(1.5), Inches(6.3), Inches(2.5), NAVY, r=0.02)
    txt(s, Inches(0.8), Inches(1.7), Inches(5.8), Inches(0.3), "Containerization", 15, TEAL, True)
    txt(s, Inches(0.8), Inches(2.1), Inches(5.8), Inches(1.6),
        "Two-stage Dockerfile -- gcc and pip in the builder, only "
        "compiled packages in runtime. Runs as non-root user "
        "(appuser) for security. Docker Compose sets CPU/memory "
        "limits and auto-restarts on failure. Healthcheck pings "
        "/health every 30 seconds.", 12, LGRAY)

    # Logging card
    rect(s, Inches(0.5), Inches(4.3), Inches(6.3), Inches(2.8), NAVY, r=0.02)
    txt(s, Inches(0.8), Inches(4.5), Inches(5.8), Inches(0.3), "Structured Logging + Testing", 15, AMBER, True)
    txt(s, Inches(0.8), Inches(4.9), Inches(5.8), Inches(2.0),
        "Every request gets a UUID. Every log entry is structured "
        "JSON. Rotating file handler caps at 10 MB with 5 backups. "
        "The pytest suite covers all endpoints: health checks, "
        "prediction classification, input validation, batch "
        "rejection, and metric accumulation. 15 tests total.", 12, LGRAY)

    # Right: code sample
    rect(s, Inches(7.0), Inches(0.3), Inches(5.8), Inches(6.8), WHITE, r=0.02)
    rect(s, Inches(7.0), Inches(0.3), Inches(5.8), Pt(5), TEAL)
    txt(s, Inches(7.3), Inches(0.5), Inches(5.2), Inches(0.3),
        "POST /predict -- Sample", 14, TEAL, True)
    txt(s, Inches(7.3), Inches(1.0), Inches(5.2), Inches(0.3), "Request:", 12, NAVY, True)
    txt(s, Inches(7.3), Inches(1.4), Inches(5.2), Inches(1.5),
        '{\n  "sepal_length": 5.1,\n  "sepal_width": 3.5,\n'
        '  "petal_length": 1.4,\n  "petal_width": 0.2\n}',
        12, SLATE, font='Consolas')
    txt(s, Inches(7.3), Inches(3.2), Inches(5.2), Inches(0.3), "Response:", 12, NAVY, True)
    txt(s, Inches(7.3), Inches(3.6), Inches(5.2), Inches(3.0),
        '{\n  "prediction": 0,\n  "class_name": "setosa",\n'
        '  "confidence": 1.0,\n  "probabilities": {\n'
        '    "setosa": 1.0,\n    "versicolor": 0.0,\n'
        '    "virginica": 0.0\n  }\n}',
        12, SLATE, font='Consolas')

    # ── SLIDE 9: WEEK 6 ───────────────────────────────────────────────────
    print("[9/10] Week 6")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_CONTENT)

    txt(s, Inches(0.9), Inches(0.3), Inches(3), Inches(0.3), "WEEK 6", 13, CORAL, True)
    badge(s, Inches(5.5), Inches(0.3), "70/100", CORAL)
    txt(s, Inches(0.9), Inches(0.7), Inches(6), Inches(0.6),
        "AI ETHICS & BIAS AUDIT", 28, NAVY, True)
    txt(s, Inches(0.9), Inches(1.4), Inches(6), Inches(0.5),
        "Audited a Logistic Regression on the UCI Adult Income "
        "dataset -- 32,561 records. Does this model treat everyone fairly?", 13, SLATE)

    # Findings (dark card)
    rect(s, Inches(0.5), Inches(2.2), Inches(6.3), Inches(2.3), NAVY, r=0.02)
    txt(s, Inches(0.8), Inches(2.4), Inches(5.8), Inches(0.3),
        "WHAT THE AUDIT FOUND", 14, CORAL, True)
    multi(s, Inches(0.8), Inches(2.8), Inches(5.8), Inches(1.5), [
        "  Disparate Impact by sex fell below the legal 80% threshold",
        "  TPR for high-income females was much lower than for males",
        "  Non-White females had the lowest positive prediction rates",
    ], 12, LGRAY)

    # Remediation (light card)
    rect(s, Inches(0.5), Inches(4.8), Inches(6.3), Inches(2.3), WHITE, r=0.02)
    rect(s, Inches(0.5), Inches(4.8), Inches(6.3), Pt(5), CORAL)
    txt(s, Inches(0.8), Inches(5.0), Inches(5.8), Inches(0.3),
        "PROPOSED FIXES", 14, CORAL, True)
    multi(s, Inches(0.8), Inches(5.4), Inches(5.8), Inches(1.5), [
        "  Resample training data to equalize representation",
        "  Apply adversarial debiasing during model training",
        "  Tune classification thresholds per demographic group",
    ], 12, SLATE)

    add_img(s, os.path.join(IMG_ETH, "fig2_disparate_impact.png"),
            Inches(7.2), Inches(0.3), width=Inches(5.6))
    add_img(s, os.path.join(IMG_ETH, "fig9_intersectional_bias.png"),
            Inches(7.2), Inches(3.8), width=Inches(5.6))

    # ── SLIDE 10: CONCLUSION ──────────────────────────────────────────────
    print("[10/10] Conclusion")
    s = prs.slides.add_slide(BL)
    set_bg_img(s, BG_TITLE)

    txt(s, Inches(0.9), Inches(0.5), Inches(6), Inches(0.3), "WRAPPING UP", 13, TEAL, True)
    txt(s, Inches(0.9), Inches(0.9), Inches(6), Inches(0.6),
        "KEY TAKEAWAYS", 32, NAVY, True)

    # Three takeaway cards
    takes = [
        ("1", "Ship it, don't just train it",
         "The deployment week scored 95 because it proved the "
         "model works in production -- not just in a notebook. "
         "A model that can't serve requests has zero value.", TEAL),
        ("2", "If you can't explain it, don't deploy it",
         "SHAP and LIME turned a black-box into something a "
         "clinician can actually review. Transparency builds "
         "the trust that adoption requires.", AMBER),
        ("3", "Fairness isn't optional",
         "The ethics audit caught real bias -- females and "
         "minorities were systematically disadvantaged. "
         "You have to check before you ship.", CORAL),
    ]
    for i, (num, title, desc, color) in enumerate(takes):
        y = Inches(1.8 + i * 1.8)
        rect(s, Inches(0.5), y, Inches(6.5), Inches(1.5), WHITE, r=0.02)
        rect(s, Inches(0.5), y, Pt(6), Inches(1.5), color)
        n = rect(s, Inches(0.8), y + Inches(0.3), Inches(0.45), Inches(0.45), color, r=0.5)
        tf = n.text_frame; p = tf.paragraphs[0]
        p.text = num; p.font.size = Pt(14); p.font.color.rgb = WHITE
        p.font.bold = True; p.font.name = 'Segoe UI'; p.alignment = PP_ALIGN.CENTER
        txt(s, Inches(1.5), y + Inches(0.15), Inches(5.2), Inches(0.35), title, 14, NAVY, True)
        txt(s, Inches(1.5), y + Inches(0.5), Inches(5.2), Inches(0.9), desc, 11, SLATE)

    # Right panel scores
    txt(s, Inches(8.5), Inches(1.8), Inches(3.8), Inches(0.4),
        "SCORES", 18, WHITE, True, PP_ALIGN.CENTER)

    scores = [("Week 1", "65", LGRAY), ("Week 2", "85", MINT),
              ("Week 3", "45", LGRAY), ("Week 4", "55", LGRAY),
              ("Week 5", "95", TEAL),  ("Week 6", "70", LGRAY)]
    for i, (wk, sc, col) in enumerate(scores):
        y = Inches(2.5 + i * 0.7)
        rect(s, Inches(8.3), y, Inches(4.0), Inches(0.55), col, r=0.1)
        txt(s, Inches(8.5), y + Inches(0.07), Inches(1.8), Inches(0.4),
            wk, 13, NAVY if col != TEAL else WHITE, True)
        txt(s, Inches(11.0), y + Inches(0.07), Inches(1.1), Inches(0.4),
            sc, 15, NAVY if col != TEAL else WHITE, True, PP_ALIGN.RIGHT)

    txt(s, Inches(8.3), Inches(6.5), Inches(4.0), Inches(0.6),
        "Thank You!", 26, WHITE, True, PP_ALIGN.CENTER)

    # ── SAVE ──────────────────────────────────────────────────────────────
    path = os.path.join(OUT, "Internship_Presentation_v5.pptx")
    prs.save(path)
    print(f"\n[OK] Saved -> {path}")
    print(f"    Slides: {len(prs.slides)}")

if __name__ == "__main__":
    build()
