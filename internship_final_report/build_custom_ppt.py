import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ── Paths ─────────────────────────────────────────────────────────────────────
BASE   = r"c:\Users\prate\Desktop\youva"
OUT    = os.path.join(BASE, "internship_final_report")

# Images
IMG_OFFICE = r"C:\Users\prate\.gemini\antigravity\brain\d283040d-239a-49c8-84f0-2e906e7e241b\office_team_1787346642152.jpg"
IMG_LAPTOP = r"C:\Users\prate\.gemini\antigravity\brain\d283040d-239a-49c8-84f0-2e906e7e241b\colleagues_laptop_1787346674433.jpg"
IMG_MEETING = r"C:\Users\prate\.gemini\antigravity\brain\d283040d-239a-49c8-84f0-2e906e7e241b\team_meeting_overhead_1787346686435.jpg"
IMG_HANDS = r"C:\Users\prate\.gemini\antigravity\brain\d283040d-239a-49c8-84f0-2e906e7e241b\typing_hands_1787346700947.jpg"

# ── Palette ───────────────────────────────────────────────────────────────────
BG       = RGBColor(0xF6, 0xF0, 0xE4)
BROWN    = RGBColor(0x4F, 0x30, 0x25)
ORANGE   = RGBColor(0xE7, 0x77, 0x32)
YELLOW   = RGBColor(0xF0, 0xB5, 0x3D)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

SW = Inches(13.333)
SH = Inches(7.5)

def _bg(slide, c=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = c

def _rect(slide, l, t, w, h, c, radius=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                l, t, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = c
    sh.line.fill.background()
    if radius: sh.adjustments[0] = radius
    return sh

def _circle(slide, l, t, d, c):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, d, d)
    sh.fill.solid(); sh.fill.fore_color.rgb = c
    sh.line.fill.background()
    return sh

def _txt(slide, l, t, w, h, text, sz=18, c=BROWN, bold=False, align=PP_ALIGN.LEFT, font='Arial'):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text; p.font.size = Pt(sz); p.font.color.rgb = c
    p.font.bold = bold; p.font.name = font; p.alignment = align
    return tb

def _img(slide, path, l, t, width=None, height=None):
    if not os.path.exists(path): return False
    kw = {}
    if width: kw['width'] = width
    if height: kw['height'] = height
    slide.shapes.add_picture(path, l, t, **kw)
    return True

def slide1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    # Decorations
    _circle(s, Inches(1), Inches(1), Inches(1), BROWN)
    _txt(s, Inches(0.8), Inches(1.8), Inches(2), Inches(2), "*", sz=100, c=BROWN, bold=True)
    _rect(s, Inches(3.5), Inches(0), Inches(2.5), Inches(1.5), ORANGE, radius=0.5)
    _circle(s, Inches(3.8), Inches(4), Inches(1.2), ORANGE)
    
    _rect(s, Inches(0), Inches(5), Inches(5), Inches(2.5), YELLOW)

    _txt(s, Inches(1), Inches(0.5), Inches(5), Inches(0.5), "AI/ML INTERNSHIP PROGRAM", sz=12, c=BROWN, bold=True)
    _txt(s, Inches(0.8), Inches(2.5), Inches(8), Inches(2), "INTERNSHIP\nREPORT", sz=65, c=BROWN, bold=True)
    _txt(s, Inches(0.8), Inches(5.2), Inches(4), Inches(1), "A Reflective Overview of My Learning\nExperience and Professional Development", sz=14, c=BROWN, bold=True)
    _txt(s, Inches(0.8), Inches(6.5), Inches(4), Inches(0.5), "Presented by: Prateek Vijay", sz=12, c=BROWN)

def slide2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(5), Inches(7.5), BROWN)
    _rect(s, Inches(7.5), Inches(0.5), Inches(5.5), Inches(6.5), YELLOW, radius=0.1)
    
    _img(s, IMG_LAPTOP, Inches(7.7), Inches(0.7), width=Inches(5.1), height=Inches(6.1))
    
    _rect(s, Inches(1), Inches(4), Inches(6.5), Inches(2), WHITE, radius=0.2)
    _txt(s, Inches(1.5), Inches(4.5), Inches(5), Inches(1), "INTRODUCTION", sz=40, c=BROWN, bold=True)
    
    _txt(s, Inches(1), Inches(0.5), Inches(3.5), Inches(3.5), 
         "My 6-week intensive AI internship was a valuable opportunity to experience real-world machine learning challenges in a professional environment.\n\nOver the course of the program, I worked closely with industry tools to develop practical skills that enhanced my academic background.\n\nThis presentation outlines the key aspects of my projects, focusing especially on AI Model Deployment and API engineering.", 
         sz=13, c=WHITE)

def slide3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _rect(s, Inches(0), Inches(5.5), Inches(13.333), Inches(2), YELLOW)
    _circle(s, Inches(-1), Inches(5), Inches(3), BROWN)
    
    _txt(s, Inches(1), Inches(1.5), Inches(6), Inches(2), "OBJECTIVES\nOF THE\nINTERNSHIP", sz=45, c=BROWN, bold=True)
    _txt(s, Inches(1), Inches(4.5), Inches(5), Inches(2), 
         "The internship was designed to bridge the gap between academic theory and AI in production reality. It offered a chance to explore how data science functions in practice and to identify the skills necessary for success in the field.", 
         sz=14, c=BROWN)
    
    _rect(s, Inches(6.5), Inches(4.2), Inches(6), Inches(0.8), ORANGE, radius=0.5)
    _txt(s, Inches(7), Inches(4.4), Inches(5), Inches(0.5), "Apply university knowledge in a real business environment", sz=14, c=WHITE)

    _rect(s, Inches(6.5), Inches(5.2), Inches(6), Inches(0.8), ORANGE, radius=0.5)
    _txt(s, Inches(7), Inches(5.4), Inches(5), Inches(0.5), "Gain insight into machine learning pipelines and workflows", sz=14, c=WHITE)

    _rect(s, Inches(6.5), Inches(6.2), Inches(6), Inches(0.8), ORANGE, radius=0.5)
    _txt(s, Inches(7), Inches(6.4), Inches(5), Inches(0.5), "Enhance professional coding, deployment and monitoring skills", sz=14, c=WHITE)

def slide4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _rect(s, Inches(0), Inches(4), Inches(13.333), Inches(3.5), YELLOW)
    
    _txt(s, Inches(1), Inches(4.5), Inches(4), Inches(2), "API\nDEPLOYMENT\nOVERVIEW", sz=45, c=BROWN, bold=True)
    _txt(s, Inches(1), Inches(6.5), Inches(5), Inches(2), 
         "The most substantial deliverable of the internship was the end-to-end deployment of the AI model. It was shipped as a production-ready REST API, focusing on resilience, monitoring, and scale.", 
         sz=12, c=BROWN)
    
    _rect(s, Inches(6), Inches(4.5), Inches(3), Inches(1.2), ORANGE, radius=0.2)
    _txt(s, Inches(6.2), Inches(4.7), Inches(2.6), Inches(0.8), "FastAPI Framework for fast, robust routing", sz=14, c=BROWN, align=PP_ALIGN.CENTER)

    _rect(s, Inches(9.5), Inches(4.5), Inches(3), Inches(1.2), ORANGE, radius=0.2)
    _txt(s, Inches(9.7), Inches(4.7), Inches(2.6), Inches(0.8), "Docker Multi-stage Builds for isolation", sz=14, c=BROWN, align=PP_ALIGN.CENTER)

    _rect(s, Inches(6), Inches(6), Inches(3), Inches(1.2), ORANGE, radius=0.2)
    _txt(s, Inches(6.2), Inches(6.2), Inches(2.6), Inches(0.8), "Thread-safe Metrics & JSON Logging", sz=14, c=BROWN, align=PP_ALIGN.CENTER)

    _rect(s, Inches(9.5), Inches(6), Inches(3), Inches(1.2), ORANGE, radius=0.2)
    _txt(s, Inches(9.7), Inches(6.2), Inches(2.6), Inches(0.8), "15 Automated Pytest tests for API logic", sz=14, c=BROWN, align=PP_ALIGN.CENTER)

def slide5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    
    _txt(s, Inches(1), Inches(1), Inches(5), Inches(1.5), "DEPLOYMENT\nTASKS AND\nRESPONSIBILITIES", sz=35, c=BROWN, bold=True)
    _txt(s, Inches(1), Inches(2.8), Inches(5), Inches(1.5), 
         "I focused on engineering the infrastructure required to serve the Random Forest classifier reliably. This meant containerizing the service and ensuring high observability.", 
         sz=12, c=BROWN)
    
    _rect(s, Inches(1), Inches(4), Inches(2.5), Inches(1), WHITE, radius=0.5)
    _txt(s, Inches(1.1), Inches(4.2), Inches(2.3), Inches(0.6), "Designed REST API endpoints using FastAPI", sz=12, c=BROWN, align=PP_ALIGN.CENTER)
    
    _rect(s, Inches(3.8), Inches(4), Inches(2.5), Inches(1), YELLOW, radius=0.5)
    _txt(s, Inches(3.9), Inches(4.2), Inches(2.3), Inches(0.6), "Containerized app with Docker Compose", sz=12, c=BROWN, align=PP_ALIGN.CENTER)

    _rect(s, Inches(1), Inches(5.2), Inches(2.5), Inches(1), YELLOW, radius=0.5)
    _txt(s, Inches(1.1), Inches(5.4), Inches(2.3), Inches(0.6), "Implemented thread-safe Prometheus metrics", sz=12, c=BROWN, align=PP_ALIGN.CENTER)

    _rect(s, Inches(3.8), Inches(5.2), Inches(2.5), Inches(1), ORANGE, radius=0.5)
    _txt(s, Inches(3.9), Inches(5.4), Inches(2.3), Inches(0.6), "Wrote extensive Pytest automation suite", sz=12, c=BROWN, align=PP_ALIGN.CENTER)

    _rect(s, Inches(7.5), Inches(0), Inches(5.833), Inches(7.5), BROWN)
    _img(s, IMG_MEETING, Inches(8), Inches(1), width=Inches(4.8), height=Inches(5.5))

def slide6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _rect(s, Inches(0), Inches(0), Inches(13.333), Inches(7.5), BROWN)
    _rect(s, Inches(1.5), Inches(1), Inches(10.333), Inches(5.5), YELLOW, radius=0.1)
    
    _txt(s, Inches(2), Inches(1.5), Inches(9.333), Inches(1), "DEPLOYMENT SKILLS GAINED", sz=45, c=BROWN, bold=True, align=PP_ALIGN.CENTER)
    _txt(s, Inches(2.5), Inches(3), Inches(8.333), Inches(3), 
         "My focus on the deployment phase allowed me to strengthen my software engineering and DevOps skills. I became proficient in building RESTful APIs using FastAPI and Pydantic validation, and learned how to containerize Python applications efficiently using multi-stage Docker builds.\n\nCrucially, I learned the importance of system observability—implementing structured JSON logging and in-process metrics to monitor latency, request volume, and model drift in a production environment.", 
         sz=16, c=BROWN, align=PP_ALIGN.CENTER)

def slide7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    _rect(s, Inches(0), Inches(5.5), Inches(13.333), Inches(2), YELLOW)
    _rect(s, Inches(0), Inches(0), Inches(5), Inches(7.5), YELLOW)
    _circle(s, Inches(-1), Inches(-1), Inches(3), WHITE)
    
    _img(s, IMG_HANDS, Inches(0.5), Inches(0.5), width=Inches(4), height=Inches(3))
    
    _txt(s, Inches(0.5), Inches(3.7), Inches(4), Inches(3.5), 
         "One of the main challenges I encountered was adjusting to the complexities of container orchestration and robust API design. Deadlines were tight, and configuring Docker health checks while keeping the image size small was tricky.\n\nI learned to manage stress under pressure, quickly adapt to new tools like Docker and FastAPI, and systematically debug container logs.", 
         sz=12, c=BROWN)
    
    _txt(s, Inches(6), Inches(6), Inches(6), Inches(1), "CHALLENGES\nFACED", sz=35, c=BROWN, bold=True)
    _circle(s, Inches(5), Inches(6), Inches(0.8), ORANGE)

def slide8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    
    _rect(s, Inches(0.5), Inches(4), Inches(5.5), Inches(2.5), ORANGE, radius=0.2)
    _txt(s, Inches(1), Inches(4.2), Inches(4.5), Inches(0.5), "KEY TAKEAWAYS", sz=25, c=WHITE, bold=True)
    _txt(s, Inches(1), Inches(4.8), Inches(4.5), Inches(1.5), 
         "Throughout my internship, I discovered the real value of resilient infrastructure in a professional AI environment. Beyond model accuracy, I learned the importance of operational stability, reproducibility, and logging in a workplace setting.", 
         sz=12, c=BROWN)

    _txt(s, Inches(7.5), Inches(4), Inches(5), Inches(0.5), "Applying theory to production helps identify real world gaps", sz=14, c=BROWN)
    _txt(s, Inches(7.5), Inches(5), Inches(5), Inches(0.5), "AI engineering requires rigorous automated API testing", sz=14, c=BROWN)
    _txt(s, Inches(7.5), Inches(6), Inches(5), Inches(0.5), "Robust observability (logs/metrics) is essential for scale", sz=14, c=BROWN)

def slide9(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    
    _txt(s, Inches(1), Inches(2), Inches(5), Inches(1), "CONCLUSION", sz=45, c=BROWN, bold=True)
    _rect(s, Inches(1), Inches(3.5), Inches(4), Inches(2.5), YELLOW, radius=0.2)
    _txt(s, Inches(1.2), Inches(3.8), Inches(3.6), Inches(2), 
         "Completing this internship was a pivotal step in my academic and professional development. It proved that I can ship a model end-to-end, earning the highest score for deployment practices.", 
         sz=12, c=BROWN)

    _rect(s, Inches(5.5), Inches(1), Inches(4.5), Inches(5.5), ORANGE, radius=0.2)
    
    _txt(s, Inches(6.5), Inches(1.5), Inches(3), Inches(0.8), "Internships are key to understanding production systems", sz=12, c=BROWN)
    _txt(s, Inches(6.5), Inches(2.8), Inches(3), Inches(0.8), "Practical exposure enhances classroom ML learning", sz=12, c=BROWN)
    _txt(s, Inches(6.5), Inches(4.1), Inches(3), Inches(0.8), "Confidence grows through active deployment and testing", sz=12, c=BROWN)
    _txt(s, Inches(6.5), Inches(5.4), Inches(3), Inches(0.8), "AI environments require containerization and monitoring", sz=12, c=BROWN)

def slide10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(s)
    
    _txt(s, Inches(1), Inches(1.5), Inches(6), Inches(2.5), "THANK YOU\nFOR YOUR\nATTENTION", sz=50, c=BROWN, bold=True)
    
    _txt(s, Inches(1.5), Inches(4.5), Inches(5), Inches(0.5), "prateek@example.com", sz=14, c=BROWN)
    _txt(s, Inches(1.5), Inches(5.2), Inches(5), Inches(0.5), "www.github.com/prateek", sz=14, c=BROWN)
    _txt(s, Inches(1.5), Inches(5.9), Inches(5), Inches(0.5), "+123-456-7890", sz=14, c=BROWN)
    
    _rect(s, Inches(7.5), Inches(0), Inches(5.833), Inches(7.5), WHITE)
    _img(s, IMG_OFFICE, Inches(8), Inches(1), width=Inches(4.8), height=Inches(5.5))

def main():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    
    slide1(prs)
    slide2(prs)
    slide3(prs)
    slide4(prs)
    slide5(prs)
    slide6(prs)
    slide7(prs)
    slide8(prs)
    slide9(prs)
    slide10(prs)
    
    prs.save(os.path.join(OUT, "Internship_Presentation_Deployed.pptx"))

if __name__ == "__main__":
    main()
